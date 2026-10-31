"""Ingestion: parse → structure → chunks → records → evidence → job record.

The fixtures build real files in memory (a hand-written PDF, a python-docx document, an
openpyxl workbook, delimited text) so the parsers are exercised for real rather than mocked.
"""

from __future__ import annotations

import datetime as dt
from itertools import pairwise

import pytest
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from tests.fixtures.documents_bytes import (
    docx_bytes,
    minimal_pdf,
    scanned_pdf,
    xlsx_bytes,
    xlsx_formula_cached_bytes,
)

from drillai.core.errors import ExtractionFailed, IngestionError, UnsupportedFormat
from drillai.db.models import (
    Base,
    Document,
    DocumentChunk,
    DocumentPage,
    DocumentRegion,
    EvidenceLink,
    ExtractedRecord,
    IngestionJob,
    Organization,
    RawArtifact,
    Well,
)
from drillai.ingestion.parsers import parse_bytes
from drillai.ingestion.pipeline import IngestionPipeline, build_chunks, normalize_text


@pytest.fixture
async def session():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(bind=engine, expire_on_commit=False)
    async with factory() as sess:
        yield sess
    await engine.dispose()


@pytest.fixture
async def context(session: AsyncSession, tmp_path):
    org = Organization(slug="ing", name="Ingestion Co")
    session.add(org)
    await session.flush()
    from drillai.db.models import Project

    project = Project(org_id=org.id, name="Development")
    session.add(project)
    await session.flush()
    well = Well(org_id=org.id, project_id=project.id, name="NF-12")
    session.add(well)
    await session.flush()
    from drillai.ingestion.storage import LocalBlobStore

    store = LocalBlobStore(tmp_path / "blobs")
    pipeline = IngestionPipeline(session, store, org_id=org.id)
    return {"org": org, "project": project, "well": well, "store": store, "pipeline": pipeline}


# --------------------------------------------------------------------------- parsers


def test_pdf_parser_extracts_pages_and_text():
    parsed = parse_bytes(minimal_pdf(), filename="ddr.pdf")
    assert parsed.parser == "pdf"
    assert len(parsed.pages) == 1
    assert "Daily Drilling Report" in parsed.pages[0].text
    assert parsed.pages[0].width_pt == pytest.approx(612)


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("1,2", 1.2),
        ("0,625", 0.625),
        ("1,200", None),
        ("1,200.5", 1200.5),
        ("1.200,5", 1200.5),
        ("1,234,567", 1_234_567.0),
        ("1.234.567", 1_234_567.0),
        ("NaN", None),
        ("Infinity", None),
        ("-Infinity", None),
        (float("nan"), None),
        (float("inf"), None),
        (True, None),
    ],
)
def test_extractor_number_parser_rejects_ambiguous_and_nonfinite_values(raw, expected):
    from drillai.ingestion.extractors import _to_float

    actual = _to_float(raw)
    if expected is None:
        assert actual is None
    else:
        assert actual == pytest.approx(expected)


def test_depth_locator_does_not_match_the_tail_of_a_grouped_number():
    from drillai.ingestion.extractors import extract

    parsed = parse_bytes(
        b"Bit depth: 1,200 ft\n", filename="ambiguous-depth.txt", content_type="text/plain"
    )
    drafts, _warnings = extract(parsed, doc_type="ddr")

    assert not [draft for draft in drafts if draft.record_type == "depth_reading"]


def test_scanned_pdf_is_flagged_not_silently_empty():
    parsed = parse_bytes(scanned_pdf(), filename="scan.pdf")
    assert parsed.metadata["ocr_required"] is True
    assert parsed.metadata["ocr_required_pages"] == [1]
    assert any("scanned" in warning for warning in parsed.warnings)


def test_mixed_text_and_scanned_pdf_marks_only_the_pages_needing_ocr():
    import io

    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(minimal_pdf())))
    writer.append(PdfReader(io.BytesIO(scanned_pdf())))
    output = io.BytesIO()
    writer.write(output)

    parsed = parse_bytes(output.getvalue(), filename="mixed-text-scan.pdf")

    assert parsed.metadata["ocr_required"] is True
    assert parsed.metadata["ocr_required_pages"] == [2]
    assert "Daily Drilling Report" in parsed.pages[0].text
    assert not parsed.pages[1].text
    assert any("OCR is required for those pages" in warning for warning in parsed.warnings)


def test_docx_and_xlsx_and_csv_parsers():
    docx_doc = parse_bytes(docx_bytes(), filename="program.docx")
    assert any(region.kind == "table" for region in docx_doc.pages[0].regions)

    xlsx_doc = parse_bytes(xlsx_bytes(), filename="mud.xlsx")
    assert xlsx_doc.pages[0].regions[0].table["rows"] == 3

    csv_doc = parse_bytes(b"md,incl,azi\n0,0,0\n500,1.5,45\n", filename="survey.csv")
    assert csv_doc.pages[0].regions[0].table["header"] == ["md", "incl", "azi"]


def test_csv_sniffs_semicolon_bom_quotes_and_preserves_physical_line_locators():
    data = (
        b"\xef\xbb\xbfTime;Mud Weight (ppg);Comment\r\n"
        b'06:00;9.2;"brine; high"\r\n'
        b'12:00;0;"line one\r\nline two"\r\n'
    )
    parsed = parse_bytes(data, filename="mud-check.csv")
    table = parsed.pages[0].regions[0].table
    assert table is not None
    assert parsed.metadata["encoding"] == "utf-8-sig"
    assert parsed.metadata["delimiter"] == ";"
    assert table["header"] == ["Time", "Mud Weight (ppg)", "Comment"]
    assert table["cells"][1][2]["text"] == "brine; high"
    assert table["cells"][2][1]["text"] == "0"
    assert table["cells"][2][2]["text"] == "line one\r\nline two"
    assert table["cells"][2][2]["source_locator"]["row"] == 3
    assert table["cells"][2][2]["source_locator"]["end_line"] == 4


def test_malformed_quoted_csv_fails_explicitly():
    with pytest.raises(ExtractionFailed, match="malformed delimited text"):
        parse_bytes(b'name,value\n"broken,12.4', filename="malformed.csv")


def test_header_detection_uses_whole_tokens_not_accidental_substrings():
    from drillai.ingestion.parsers import detect_header_row

    index, header, confidence = detect_header_row(
        [
            ["Fluid", "Recommended", "Valid", "Grid"],
            ["Well", "Date", "Time", "Value"],
        ]
    )

    assert index == 1
    assert header == ["Well", "Date", "Time", "Value"]
    assert confidence == 1.0


def test_operation_extractor_does_not_treat_decoder_as_an_operation_code_header():
    from drillai.ingestion.extractors import extract

    parsed = parse_bytes(b"Decoder,Value\nflow-rate decoder,12.4\n", filename="codes.csv")
    drafts, warnings = extract(parsed, doc_type="ddr")

    assert warnings == []
    assert not any(draft.record_type == "operation_row" for draft in drafts)


def test_delimited_and_xlsx_processing_limits_are_enforced(monkeypatch):
    from drillai.ingestion import parsers

    monkeypatch.setattr(parsers, "MAX_SPREADSHEET_ROWS", 1)
    with pytest.raises(ExtractionFailed, match="row processing limit"):
        parse_bytes(b"header\nvalue", filename="rows.csv")

    monkeypatch.setattr(parsers, "MAX_SPREADSHEET_ROWS", 100_000)
    monkeypatch.setattr(parsers, "MAX_SPREADSHEET_COLUMNS", 1)
    with pytest.raises(ExtractionFailed, match="column processing limit"):
        parse_bytes(b"a,b\n1,2", filename="columns.csv")

    monkeypatch.setattr(parsers, "MAX_SPREADSHEET_COLUMNS", 2_000)
    monkeypatch.setattr(parsers, "MAX_SPREADSHEET_CELLS", 2)
    with pytest.raises(ExtractionFailed, match="cell processing limit"):
        parse_bytes(b"a,b\n1,2", filename="cells.csv")

    monkeypatch.setattr(parsers, "MAX_XLSX_UNCOMPRESSED_BYTES", 1)
    with pytest.raises(ExtractionFailed, match="uncompressed processing limit"):
        parse_bytes(xlsx_bytes(), filename="oversized.xlsx")


def test_xlsx_rejects_worksheet_count_over_the_processing_limit(monkeypatch):
    import io

    import openpyxl

    from drillai.ingestion import parsers

    workbook = openpyxl.Workbook()
    workbook.create_sheet("Extra")
    data = io.BytesIO()
    workbook.save(data)
    workbook.close()
    monkeypatch.setattr(parsers, "MAX_PAGES", 1)

    with pytest.raises(ExtractionFailed, match="worksheet processing limit"):
        parse_bytes(data.getvalue(), filename="too-many-sheets.xlsx")


def test_xlsx_sparse_oversized_dimensions_are_rejected_before_openpyxl_load(monkeypatch):
    import io
    import zipfile

    import openpyxl
    from tests.fixtures.documents_bytes import xlsx_bytes

    source = xlsx_bytes()
    malformed = io.BytesIO()
    replaced = False
    with zipfile.ZipFile(io.BytesIO(source)) as original, zipfile.ZipFile(malformed, "w") as rebuilt:
        for member in original.infolist():
            payload = original.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                old_dimension = b'<dimension ref="A1:E3"/>'
                assert old_dimension in payload
                payload = payload.replace(old_dimension, b'<dimension ref="A1:XFD1048576"/>', 1)
                replaced = True
            rebuilt.writestr(member, payload)
    assert replaced
    monkeypatch.setattr(
        openpyxl,
        "load_workbook",
        lambda *_args, **_kwargs: pytest.fail("openpyxl must not load oversized worksheet dimensions"),
    )

    with pytest.raises(ExtractionFailed, match="row limit"):
        parse_bytes(malformed.getvalue(), filename="sparse-oversized.xlsx")


@pytest.mark.parametrize(
    ("filename", "source", "member_name", "root_tag"),
    [
        ("unsafe.docx", docx_bytes, "word/document.xml", b"<w:document"),
        ("unsafe.xlsx", xlsx_bytes, "xl/worksheets/sheet1.xml", b"<worksheet"),
    ],
)
@pytest.mark.parametrize(
    "declaration",
    [
        b'<!DOCTYPE worksheet [<!ENTITY payload "expanded">]>',
        b'<!DOCTYPE worksheet SYSTEM "file:///etc/passwd">',
    ],
)
def test_office_xml_preflight_rejects_dtd_and_entity_declarations(
    filename, source, member_name, root_tag, declaration
):
    import io
    import zipfile

    original = source()
    unsafe = io.BytesIO()
    replaced = False
    with zipfile.ZipFile(io.BytesIO(original)) as source_archive, zipfile.ZipFile(
        unsafe, "w", compression=zipfile.ZIP_DEFLATED
    ) as target_archive:
        for member in source_archive.infolist():
            payload = source_archive.read(member.filename)
            if member.filename == member_name:
                assert root_tag in payload
                payload = payload.replace(root_tag, declaration + root_tag, 1)
                replaced = True
            target_archive.writestr(member, payload)
    assert replaced

    with pytest.raises(ExtractionFailed, match="forbidden DTD/entity"):
        parse_bytes(unsafe.getvalue(), filename=filename)


def test_csv_and_docx_expansion_limits_are_enforced_before_materialization(monkeypatch):
    import io
    import zipfile

    from drillai.ingestion import parsers

    monkeypatch.setattr(parsers, "MAX_SPREADSHEET_RECORD_CHARS", 12)
    with pytest.raises(ExtractionFailed, match="record exceeds"):
        parse_bytes(b"header\n" + b"x" * 20, filename="oversized-record.csv")

    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as zipped:
        zipped.writestr("word/document.xml", b"x" * 200)
    monkeypatch.setattr(parsers, "MAX_DOCX_UNCOMPRESSED_BYTES", 100)
    with pytest.raises(ExtractionFailed, match="uncompressed processing limit"):
        parsers._validate_docx_archive(archive.getvalue())


def test_mud_table_extraction_uses_original_xlsx_cells():
    from drillai.ingestion.extractors import extract

    parsed = parse_bytes(xlsx_bytes(), filename="mud-check.xlsx")
    drafts, _warnings = extract(parsed, doc_type="other")
    mud = next(draft for draft in drafts if draft.record_type == "mud_properties")

    assert mud.payload["mud_weight"] == 9.2
    assert mud.payload["plastic_viscosity"] == 18.0
    assert mud.field_locations["mud_weight"] == {
        "format": "xlsx",
        "row": 2,
        "column": 2,
        "sheet": "Mud Check",
        "cell": "B2",
        "region_order": 0,
    }
    assert mud.field_locations["plastic_viscosity"]["cell"] == "D2"
    assert mud.region_index == 0


def test_operation_duration_zero_and_explicit_units_keep_source_provenance():
    from drillai.ingestion.extractors import extract

    parsed = parse_bytes(
        b"Operation,Duration (min),Start Time\n"
        b"zero minutes,0,06:00\n"
        b"ninety minutes,90,07:00\n"
        b"zero clock,00:00,08:00\n",
        filename="durations.csv",
    )
    drafts, warnings = extract(parsed, doc_type="ddr")
    operations = [draft for draft in drafts if draft.record_type == "operation_row"]

    assert warnings == []
    assert [draft.payload["duration_hours"] for draft in operations] == [0.0, 1.5, 0.0]
    assert [draft.unit_context["duration_hours"] for draft in operations] == ["min", "min", "min"]
    assert [
        parsed.pages[0].regions[0].table["cells"][int(draft.payload["row_index"]) - 1][1]["source_locator"]["row"]
        for draft in operations
    ] == [2, 3, 4]
    assert [draft.payload["start time"] for draft in operations] == ["06:00", "07:00", "08:00"]
    assert [draft.payload["duration_hours"] for draft in operations] == [0.0, 1.5, 0.0]


def test_operation_duration_bound_is_applied_after_unit_conversion():
    from drillai.ingestion.extractors import OperationTableExtractor

    parse_duration = OperationTableExtractor._duration_hours
    assert parse_duration({"duration (minutes)": "120"}) == pytest.approx(2.0)
    assert parse_duration({"duration (days)": "4"}) == pytest.approx(96.0)
    assert parse_duration({"duration (days)": "5"}) is None
    assert parse_duration({"duration (hours)": "100"}) is None


def test_operation_clock_duration_rejects_invalid_minute_component():
    from drillai.ingestion.extractors import extract

    parsed = parse_bytes(
        b"Operation,Duration\ninvalid clock,01:99\n",
        filename="invalid-duration.csv",
    )
    drafts, _warnings = extract(parsed, doc_type="ddr")
    operation = next(draft for draft in drafts if draft.record_type == "operation_row")

    assert "duration_hours" not in operation.payload


def test_operation_depth_uses_the_declared_depth_cell_not_a_comment_first():
    from drillai.ingestion.extractors import extract

    parsed = parse_bytes(
        b"Operation,Measured Depth (m),Description\n"
        b"Drilling ahead,2410,Drilled from 2350 m to 2410 m\n",
        filename="operation-depth.csv",
    )
    drafts, warnings = extract(parsed, doc_type="ddr")
    operation = next(draft for draft in drafts if draft.record_type == "operation_row")

    assert warnings == []
    assert operation.payload["depth_md_si"] == pytest.approx(2410.0)
    assert operation.unit_context["depth_md_si"] == "m"
    assert operation.field_locations["depth_md_si"]["row"] == 2
    assert operation.field_locations["depth_md_si"]["column"] == 2


def test_generic_depth_column_is_not_mislabeled_as_measured_depth():
    from drillai.ingestion.extractors import extract

    parsed = parse_bytes(
        b"Operation,Depth (m)\nDrilling ahead,2410\n",
        filename="generic-depth.csv",
    )
    drafts, _warnings = extract(parsed, doc_type="ddr")
    operation = next(draft for draft in drafts if draft.record_type == "operation_row")

    assert operation.payload["depth (m)"] == "2410"
    assert "depth_md_si" not in operation.payload


def test_operation_description_depth_is_not_promoted_without_an_explicit_depth_column():
    from drillai.ingestion.extractors import extract

    parsed = parse_bytes(
        b"Operation,Description\nDrilling ahead,Drilled from 2350 m to 2410 m\n",
        filename="operation-no-depth-column.csv",
    )
    drafts, _warnings = extract(parsed, doc_type="ddr")
    operation = next(draft for draft in drafts if draft.record_type == "operation_row")

    assert "depth_md_si" not in operation.payload
    assert "depth_md_si" not in operation.field_locations


def test_survey_header_matching_does_not_promote_vertical_depth_as_measured_depth():
    from drillai.ingestion.extractors import extract

    parsed = parse_bytes(
        b"Vertical Depth (m),Inclination (deg),Azimuth (deg)\n500,12,180\n",
        filename="survey-tvd-only.csv",
    )
    drafts, _warnings = extract(parsed, doc_type="survey_report")

    assert not any(draft.record_type == "survey_station" for draft in drafts)


def test_survey_converts_declared_depth_and_angle_units_and_keeps_raw_unit_context():
    from drillai.ingestion.extractors import extract

    parsed = parse_bytes(
        b"MD (ft),Inclination (rad),Azimuth (rad)\n500,1.5707963267948966,3.141592653589793\n",
        filename="survey-units.csv",
    )
    drafts, warnings = extract(parsed, doc_type="survey_report")
    station = next(draft for draft in drafts if draft.record_type == "survey_station")

    assert warnings == []
    assert station.payload["md_si"] == pytest.approx(152.4)
    assert station.payload["inclination_deg"] == pytest.approx(90.0)
    assert station.payload["azimuth_deg"] == pytest.approx(180.0)
    assert station.unit_context == {"md_si": "ft", "inclination_deg": "rad", "azimuth_deg": "rad"}
    assert station.field_locations["md_si"]["row"] == 2
    assert station.field_locations["inclination_deg"]["column"] == 2
    assert station.field_locations["azimuth_deg"]["column"] == 3


def test_survey_does_not_coerce_ambiguous_depth_or_nonfinite_angles():
    from drillai.ingestion.extractors import extract

    parsed = parse_bytes(
        b'MD (m),Inclination (deg),Azimuth (deg)\n"1,200",NaN,Infinity\n',
        filename="survey-invalid-numeric.csv",
    )
    drafts, _warnings = extract(parsed, doc_type="survey_report")

    assert not [draft for draft in drafts if draft.record_type == "survey_station"]

    valid_depth = parse_bytes(
        b"MD (m),Inclination (deg),Azimuth (deg)\n500,NaN,Infinity\n",
        filename="survey-nonfinite-angles.csv",
    )
    drafts, _warnings = extract(valid_depth, doc_type="survey_report")
    station = next(draft for draft in drafts if draft.record_type == "survey_station")
    assert station.payload == {"md_si": 500.0}
    assert "inclination_deg" not in station.field_locations
    assert "azimuth_deg" not in station.field_locations
    assert "ambiguous_or_invalid_inclination" in station.quality_flags
    assert "ambiguous_or_invalid_azimuth" in station.quality_flags


@pytest.mark.parametrize(
    ("header", "warning_fragment"),
    [
        ("MD (km),Inclination (deg),Azimuth (deg)", "unsupported measured-depth unit"),
        ("MD (m),Inclination (grad),Azimuth (deg)", "unsupported survey angle unit"),
    ],
)
def test_survey_rejects_explicit_but_unsupported_units(header: str, warning_fragment: str):
    from drillai.ingestion.extractors import extract

    parsed = parse_bytes(
        f"{header}\n1,90,180\n".encode(),
        filename="survey-unsupported-units.csv",
    )
    drafts, warnings = extract(parsed, doc_type="survey_report")

    assert not any(draft.record_type == "survey_station" for draft in drafts)
    assert any(warning_fragment in warning for warning in warnings)


def test_extraction_deduplication_preserves_equal_values_from_distinct_source_locations():
    from dataclasses import replace

    from drillai.ingestion.extractors import ExtractedRecordDraft, dedupe

    first = ExtractedRecordDraft(
        record_type="mud_properties",
        payload={"mud_weight": 9.2},
        page_number=1,
        region_index=0,
        method="label_value",
        method_version="1.0.0",
        excerpt="Mud weight: 9.2 ppg",
        field_locations={"mud_weight": {"format": "csv", "row": 2, "column": 2}},
    )
    same_source = replace(first)
    other_source = replace(
        first,
        field_locations={"mud_weight": {"format": "csv", "row": 3, "column": 2}},
        excerpt="Mud weight: 9.2 ppg",
    )

    assert dedupe([first, same_source]) == [first]
    assert dedupe([first, other_source]) == [first, other_source]


def test_mud_text_range_is_not_promoted_but_single_reading_retains_its_locator():
    from drillai.ingestion.extractors import extract

    range_document = parse_bytes(
        b"Mud weight ranged from 9.2 to 9.4 ppg during the shift.\n",
        filename="mud-range.txt",
        content_type="text/plain",
    )
    range_drafts, _warnings = extract(range_document, doc_type="mud_report")
    assert not [draft for draft in range_drafts if draft.record_type == "mud_properties"]

    single_document = parse_bytes(
        b"Mud weight: 9.2 ppg\n", filename="mud-single.txt", content_type="text/plain"
    )
    single_drafts, _warnings = extract(single_document, doc_type="mud_report")
    mud = next(draft for draft in single_drafts if draft.record_type == "mud_properties")
    locator = mud.field_locations["mud_weight"]
    assert mud.payload["mud_weight"] == pytest.approx(9.2)
    assert mud.unit_context["mud_weight"] == "ppg"
    assert single_document.pages[0].text[locator["char_start"]:locator["char_end"]] == "Mud weight: 9.2 ppg"


def test_docx_locator_uses_the_paragraph_not_identical_table_text():
    import io

    import docx

    from drillai.ingestion.extractors import extract

    source_text = "Mud weight: 9.2 ppg"
    source = docx.Document()
    source.add_table(rows=1, cols=1).cell(0, 0).text = source_text
    source.add_paragraph(source_text)
    output = io.BytesIO()
    source.save(output)

    parsed = parse_bytes(
        output.getvalue(),
        filename="duplicate-table-paragraph.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    drafts, _warnings = extract(parsed, doc_type="other")
    mud = next(draft for draft in drafts if draft.record_type == "mud_properties")
    locator = mud.field_locations["mud_weight"]
    paragraph_region = next(
        region
        for region in parsed.pages[0].regions
        if region.kind == "text" and region.text == source_text
    )

    assert locator["kind"] == "text_line"
    assert locator["char_start"] == parsed.pages[0].text.rfind(source_text)
    assert parsed.pages[0].text[locator["char_start"]:locator["char_end"]] == source_text
    assert locator["region_order"] == paragraph_region.order_index


def test_report_date_locator_tracks_the_valid_date_that_was_parsed():
    from drillai.ingestion.extractors import extract

    parsed = parse_bytes(
        b"Report date: 2026-02-30\nIssued on: 2026-03-15\n",
        filename="date-locator.txt",
        content_type="text/plain",
    )
    drafts, _warnings = extract(parsed, doc_type="ddr")
    header = next(draft for draft in drafts if draft.record_type == "report_header")
    locator = header.field_locations["report_date"]

    assert header.payload["report_date"] == "2026-03-15"
    assert parsed.pages[0].text[locator["char_start"]:locator["char_end"]] == "2026-03-15"


def test_csv_header_fields_keep_original_cell_locators():
    from tests.fixtures.synthetic_ddr import SYNTHETIC_DDR_CSV

    from drillai.ingestion.extractors import extract

    parsed = parse_bytes(SYNTHETIC_DDR_CSV, filename="header.csv")
    drafts, _warnings = extract(parsed, doc_type="ddr")
    header = next(draft for draft in drafts if draft.record_type == "report_header")

    assert header.payload["well_name_text"] == "NF-12"
    assert header.payload["report_date"] == "2026-03-15"
    assert header.region_index == 0
    well_locator = header.field_locations["well_name_text"]
    date_locator = header.field_locations["report_date"]
    assert (well_locator["format"], well_locator["row"], well_locator["column"]) == ("csv", 4, 2)
    assert (date_locator["format"], date_locator["row"], date_locator["column"]) == ("csv", 3, 2)


def test_xlsx_preserves_formula_cache_types_hidden_rows_and_source_cells():
    parsed = parse_bytes(xlsx_formula_cached_bytes(), filename="operations.xlsx")
    page = parsed.pages[0]
    table = page.regions[0].table
    assert table is not None
    assert table["header_row"] == 2
    formula = table["cells"][3][1]
    assert formula["formula"] == "=6+6"
    assert formula["cached_value"] == 12
    assert formula["cache_available"] is True
    assert formula["text"] == "" and formula["value_type"] == "formula"
    assert formula["source_locator"]["cell"] == "B4"
    assert formula["source_locator"]["row"] == 4
    zero = table["cells"][4][1]
    blank = table["cells"][6][1]
    assert zero["text"] == "0" and zero["value"] == 0 and zero["value_type"] == "number"
    assert blank["text"] == "" and blank["value"] is None and blank["value_type"] == "blank"
    hidden = table["cells"][5][0]
    assert hidden["text"] == "" and hidden["display_text"] == "Hidden operation"
    assert hidden["hidden_row"] is True
    assert "Hidden operation" not in page.text and "99.9" not in parsed.pages[1].text
    assert "=6+6" not in page.text and "12" not in page.text
    assert "A1:C1" in table["merged_ranges"]
    assert any("formula cell" in warning for warning in parsed.warnings)
    assert any("hidden row" in warning for warning in parsed.warnings)
    assert parsed.pages[1].layout["sheet_state"] == "hidden"

    from drillai.ingestion.extractors import extract

    drafts, _warnings = extract(parsed, doc_type="ddr")
    operation_rows = [draft for draft in drafts if draft.record_type == "operation_row"]
    assert [draft.payload["row_index"] for draft in operation_rows] == [4, 5, 7]
    assert "duration (h)" not in operation_rows[0].payload  # cached formula is not silently treated as a measurement
    assert operation_rows[1].payload["duration (h)"] == "0"
    assert all(draft.payload.get("activity") != "Hidden operation" for draft in operation_rows)


def test_unsupported_format_is_rejected_explicitly():
    with pytest.raises(UnsupportedFormat):
        parse_bytes(b"\x00\x01binary", filename="mystery.bin", content_type="application/octet-stream")


def test_chunking_respects_boundaries_and_overlap():
    text = "\n\n".join(f"Paragraph {index} " + "word " * 160 for index in range(4))
    chunks = build_chunks(text, page_number=1, region_id=None)
    assert len(chunks) >= 2
    assert all(len(chunk.text) >= 80 for chunk in chunks)
    assert all(chunk.char_start >= 0 for chunk in chunks)
    assert normalize_text("a   b\t\tc") == "a b c"


def test_chunking_bounds_a_single_unbroken_paragraph_and_preserves_offsets():
    text = "x" * 5000
    chunks = build_chunks(
        text,
        page_number=1,
        region_id="reg_test",
        target_chars=1200,
        overlap_chars=200,
    )

    assert len(chunks) == 5
    assert all(len(chunk.text) <= 1200 for chunk in chunks)
    assert all(chunk.region_id == "reg_test" for chunk in chunks)
    assert all(text[chunk.char_start:chunk.char_end] == chunk.text for chunk in chunks)
    for previous, current in pairwise(chunks):
        assert current.text.startswith(previous.text[-200:])


def test_chunking_rejects_invalid_bounds():
    with pytest.raises(ValueError, match="target_chars"):
        build_chunks("x" * 100, page_number=1, region_id=None, target_chars=79)
    with pytest.raises(ValueError, match="overlap_chars"):
        build_chunks("x" * 100, page_number=1, region_id=None, overlap_chars=1200)
    with pytest.raises(ValueError, match="depth_from_si"):
        build_chunks(
            "x" * 100,
            page_number=1,
            region_id=None,
            depth_from_si=10**10000,
        )


# --------------------------------------------------------------------------- pipeline


async def test_pdf_ingestion_creates_full_provenance_chain(session, context):
    pipeline = context["pipeline"]
    outcome = await pipeline.ingest(
        minimal_pdf(),
        filename="ddr-2026-01-05.pdf",
        content_type="application/pdf",
        project_id=context["project"].id,
        well_id=context["well"].id,
        triggered_by="usr_1",
    )
    document = outcome.document
    assert document is not None
    assert outcome.job.status == "succeeded"
    # `extracted` is the canonical name for "the machine has read this and nothing was reported
    # incomplete". It replaced `ingested`, which appeared in no vocabulary — see
    # `documents/lifecycle.py` for the whole mapping.
    assert document.status == "extracted"
    assert document.doc_type == "ddr"
    assert document.well_name_text == "NF-12"
    assert document.rig_name_text == "Rig 42"
    assert document.page_count == 1
    assert outcome.stats["regions"] >= 1
    assert outcome.chunk_count >= 1

    job = (await session.execute(select(IngestionJob))).scalars().one()
    assert job.id == outcome.job.id
    assert job.document_id == document.id
    assert job.duration_ms is not None
    assert job.pipeline["extractors"]["report_header"] == "1.1.0"

    pages = (await session.execute(select(DocumentPage))).scalars().all()
    assert [page.page_number for page in pages] == [1]

    chunks = (await session.execute(select(DocumentChunk))).scalars().all()
    assert all(chunk.well_id == context["well"].id for chunk in chunks)
    assert all(chunk.doc_type == "ddr" for chunk in chunks)

    records = (await session.execute(select(ExtractedRecord))).scalars().all()
    types = {record.record_type for record in records}
    assert "report_header" in types
    assert "mud_properties" in types
    assert "depth_reading" in types
    mud = next(record for record in records if record.record_type == "mud_properties")
    assert mud.payload["mud_weight"] == pytest.approx(9.2)
    assert mud.unit_context.get("mud_weight") == "ppg"
    # A machine wrote this and nothing has checked it yet: `extracted`, not `unvalidated` (which was
    # never a declared state) and not `validated` (which would claim a check that did not happen).
    assert mud.validation_state == "extracted"
    depth = next(record for record in records if record.record_type == "depth_reading")
    assert depth.payload["bit_depth_si"] == pytest.approx(1500.0)

    links = (await session.execute(select(EvidenceLink))).scalars().all()
    assert len(links) == len(records)
    assert all(link.document_id == document.id for link in links)
    assert all(link.locator.get("page") == 1 for link in links)


async def test_page_level_evidence_links_resolve_to_a_real_page(session, context, monkeypatch):
    """A document_page edge must point to the DocumentPage row, not the raw-blob id."""
    from drillai.ingestion.extractors import ExtractedRecordDraft

    def page_only_extractor(parsed, *, doc_type):
        return (
            [
                ExtractedRecordDraft(
                    record_type="page_probe",
                    payload={"value": 12.4},
                    page_number=1,
                    excerpt="Mud Weight: 9.2 ppg",
                    source_locator={"format": "pdf", "kind": "page_text", "page": 1},
                )
            ],
            [],
        )

    monkeypatch.setattr("drillai.ingestion.pipeline.extract", page_only_extractor)
    outcome = await context["pipeline"].ingest(
        minimal_pdf(),
        filename="page-locators.pdf",
        content_type="application/pdf",
        project_id=context["project"].id,
        well_id=context["well"].id,
    )
    document = outcome.document
    assert document is not None
    page = (await session.execute(select(DocumentPage).where(DocumentPage.document_id == document.id))).scalars().one()
    links = (
        await session.execute(select(EvidenceLink).where(EvidenceLink.document_id == document.id))
    ).scalars().all()
    page_links = [link for link in links if link.evidence_kind == "document_page"]
    assert page_links, "the fixture must exercise the page-level evidence path"
    assert all(link.evidence_id == page.id for link in page_links)
    assert all(link.locator["precision"] == "page" for link in page_links)


async def test_identical_upload_is_deduplicated(session, context):
    pipeline = context["pipeline"]
    first = await pipeline.ingest(minimal_pdf(), filename="ddr.pdf", project_id=context["project"].id)
    second = await pipeline.ingest(minimal_pdf(), filename="ddr-copy.pdf", project_id=context["project"].id)
    assert second.reused_existing is True
    assert second.job.status == "skipped_duplicate"
    assert second.document is not None and second.document.id == first.document.id
    assert len((await session.execute(select(RawArtifact))).scalars().all()) == 1
    assert len((await session.execute(select(Document))).scalars().all()) == 1


@pytest.mark.parametrize(
    ("message", "raw_conflict", "document_conflict"),
    [
        ("UNIQUE constraint failed: raw_artifacts.org_id, raw_artifacts.sha256", True, False),
        ("UNIQUE constraint failed: documents.org_id, documents.logical_key", False, True),
        ("UNIQUE constraint failed: ingestion_jobs.org_id, ingestion_jobs.id", False, False),
        ("CHECK constraint failed: document_status", False, False),
    ],
)
def test_only_known_ingestion_identity_conflicts_are_recoverable(
    message, raw_conflict, document_conflict
):
    from sqlalchemy.exc import IntegrityError

    from drillai.ingestion.pipeline import (
        _is_logical_document_unique_conflict,
        _is_raw_artifact_unique_conflict,
    )

    error = IntegrityError("INSERT", {}, Exception(message))
    assert _is_raw_artifact_unique_conflict(error) is raw_conflict
    assert _is_logical_document_unique_conflict(error) is document_conflict


async def test_unrelated_integrity_error_is_not_recovered_as_a_upload_race(context, monkeypatch):
    from sqlalchemy.exc import IntegrityError

    pipeline = context["pipeline"]
    failure = IntegrityError("INSERT", {}, Exception("CHECK constraint failed: unrelated_status"))

    async def fail_ingest_inner(data, **kwargs):
        raise failure

    async def fail_if_recovery_is_attempted(*args, **kwargs):
        pytest.fail("unrelated integrity failures must not enter duplicate-upload recovery")

    monkeypatch.setattr(pipeline, "_ingest_inner", fail_ingest_inner)
    monkeypatch.setattr(pipeline, "_recover_duplicate_upload", fail_if_recovery_is_attempted)
    with pytest.raises(IntegrityError, match="unrelated_status"):
        await pipeline.ingest(
            b"arbitrary integrity failure probe",
            filename="probe.txt",
            content_type="text/plain",
            project_id=context["project"].id,
            commit=True,
        )


async def test_csv_survey_rows_become_records_with_depths(session, context):
    csv = b"MD (m),Inclination (deg),Azimuth (deg)\n0,0.0,0.0\n500,1.5,45.0\n1000,3.0,90.0\n"
    outcome = await context["pipeline"].ingest(
        csv,
        filename="survey.csv",
        content_type="text/csv",
        well_id=context["well"].id,
        project_id=context["project"].id,
    )
    records = (
        await session.execute(select(ExtractedRecord).where(ExtractedRecord.record_type == "survey_station"))
    ).scalars().all()
    assert [record.depth_md_si for record in records] == [500.0, 1000.0]
    assert records[0].payload["inclination_deg"] == pytest.approx(1.5)
    assert outcome.stats["records"] == len(outcome.record_ids)

    assert outcome.document is not None
    chunks = (
        await session.execute(select(DocumentChunk).where(DocumentChunk.document_id == outcome.document.id))
    ).scalars().all()
    assert chunks
    assert all(chunk.depth_from_si == pytest.approx(500.0) for chunk in chunks)
    assert all(chunk.depth_to_si == pytest.approx(1000.0) for chunk in chunks)


async def test_generic_parameter_depth_does_not_become_a_measured_depth_filter(
    session, context, monkeypatch
):
    from drillai.ingestion.extractors import ExtractedRecordDraft

    def parameter_set_with_generic_depth(parsed, *, doc_type):
        del doc_type
        region = parsed.pages[0].regions[0]
        return (
            [
                ExtractedRecordDraft(
                    record_type="parameter_set",
                    payload={"depth": 2412.0, "depth_unit": "ft"},
                    page_number=1,
                    region_index=region.order_index,
                    excerpt=(region.text or "")[:120],
                )
            ],
            [],
        )

    monkeypatch.setattr("drillai.ingestion.pipeline.extract", parameter_set_with_generic_depth)
    outcome = await context["pipeline"].ingest(
        minimal_pdf(),
        filename="generic-depth.pdf",
        content_type="application/pdf",
        well_id=context["well"].id,
        project_id=context["project"].id,
    )

    assert outcome.document is not None
    chunks = (
        await session.execute(select(DocumentChunk).where(DocumentChunk.document_id == outcome.document.id))
    ).scalars().all()
    assert chunks
    assert all(chunk.depth_from_si is None and chunk.depth_to_si is None for chunk in chunks)


async def test_report_header_dates_are_merged_before_chunk_and_record_persistence(session, context):
    outcome = await context["pipeline"].ingest(
        minimal_pdf(
            [
                "Daily Drilling Report",
                "Well: NF-12",
                "Report date: 2026-03-15",
                "Mud Weight: 9.2 ppg",
                "Additional drilling report text that ensures a retrievable region is created.",
            ]
        ),
        filename="dated-report.pdf",
        content_type="application/pdf",
        well_id=context["well"].id,
        project_id=context["project"].id,
    )

    assert outcome.document is not None
    expected_start = dt.datetime(2026, 3, 15, tzinfo=dt.UTC)
    expected_end = dt.datetime(2026, 3, 16, tzinfo=dt.UTC)
    assert outcome.document.issue_date == expected_start
    assert outcome.document.period_start == expected_start
    assert outcome.document.period_end == expected_end
    chunks = (
        await session.execute(select(DocumentChunk).where(DocumentChunk.document_id == outcome.document.id))
    ).scalars().all()
    assert chunks
    assert all(chunk.period_start == expected_start for chunk in chunks)
    assert all(chunk.period_end == expected_end for chunk in chunks)
    header = (
        await session.execute(
            select(ExtractedRecord).where(
                ExtractedRecord.document_id == outcome.document.id,
                ExtractedRecord.record_type == "report_header",
            )
        )
    ).scalars().one()
    assert header.observed_at == expected_start


async def test_explicit_partial_period_is_not_completed_from_report_date(session, context):
    explicit_start = dt.datetime(2026, 3, 1, 3, 30, tzinfo=dt.timezone(dt.timedelta(hours=5, minutes=30)))
    explicit_end = dt.datetime(2026, 3, 20, tzinfo=dt.UTC)
    period_cases = ((explicit_start, None), (None, explicit_end))

    for index, (period_start, period_end) in enumerate(period_cases):
        outcome = await context["pipeline"].ingest(
            minimal_pdf(
                [
                    "Daily Drilling Report",
                    "Well: NF-12",
                    "Report date: 2026-03-15",
                    f"Distinct period test {index}",
                    "Additional report content for temporal metadata testing.",
                ]
            ),
            filename=f"partial-period-{index}.pdf",
            content_type="application/pdf",
            well_id=context["well"].id,
            project_id=context["project"].id,
            period_start=period_start,
            period_end=period_end,
        )

        assert outcome.document is not None
        assert outcome.document.issue_date == dt.datetime(2026, 3, 15, tzinfo=dt.UTC)
        expected_start = (
            explicit_start.astimezone(dt.UTC) if period_start is not None else None
        )
        assert outcome.document.period_start == expected_start
        assert outcome.document.period_end == period_end
        chunks = (
            await session.execute(select(DocumentChunk).where(DocumentChunk.document_id == outcome.document.id))
        ).scalars().all()
        assert chunks
        assert all(chunk.period_start == expected_start for chunk in chunks)
        assert all(chunk.period_end == period_end for chunk in chunks)


async def test_csv_cell_locators_survive_extraction_and_evidence_persistence(session, context):
    csv_bytes = (
        b"MD (m),Inclination (deg),Azimuth (deg)\n"
        b"0,0.0,0.0\n"
        b"500,1.5,45.0\n"
    )
    outcome = await context["pipeline"].ingest(
        csv_bytes,
        filename="survey-locator.csv",
        content_type="text/csv",
        well_id=context["well"].id,
        project_id=context["project"].id,
    )
    record = (
        await session.execute(
            select(ExtractedRecord).where(ExtractedRecord.record_type == "survey_station")
        )
    ).scalars().one()
    link = (
        await session.execute(select(EvidenceLink).where(EvidenceLink.subject_id == record.id))
    ).scalars().one()
    assert outcome.document is not None
    assert record.region_id is not None
    assert link.locator["precision"] == "cell"
    assert link.locator["field_locations"]["md_si"]["format"] == "csv"
    assert link.locator["field_locations"]["md_si"]["row"] == 3
    assert link.locator["field_locations"]["md_si"]["column"] == 1
    assert link.quote_verified is True
    assert link.quote_check == "region_table"


async def test_docx_program_table_is_ingested(session, context):
    outcome = await context["pipeline"].ingest(
        docx_bytes(),
        filename="program.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        well_id=context["well"].id,
        project_id=context["project"].id,
    )
    assert outcome.document is not None
    assert outcome.document.has_tables is True
    records = (
        await session.execute(select(ExtractedRecord).where(ExtractedRecord.record_type == "survey_station"))
    ).scalars().all()
    assert {record.payload["md_si"] for record in records} == {500.0}

    header = (
        await session.execute(select(ExtractedRecord).where(ExtractedRecord.record_type == "report_header"))
    ).scalars().one()
    assert header.region_id is None and header.region_unknown is True
    link = (
        await session.execute(select(EvidenceLink).where(EvidenceLink.subject_id == header.id))
    ).scalars().one()
    page = (
        await session.execute(select(DocumentPage).where(DocumentPage.document_id == header.document_id))
    ).scalars().one()
    assert link.evidence_kind == "document_page" and link.evidence_id == page.id
    regions = (
        await session.execute(select(DocumentRegion).where(DocumentRegion.document_id == header.document_id))
    ).scalars().all()
    region_by_order = {region.order_index: region for region in regions}
    for field_name in ("well_name_text", "rig_name_text"):
        locator = link.locator["field_locations"][field_name]
        source_region = region_by_order[locator["region_order"]]
        assert header.payload[field_name] in (source_region.text or "")


async def test_text_locators_respect_crlf_and_indentation(session, context):
    payload = b"\r\n  Bit Depth: 1500 m\r\n"
    outcome = await context["pipeline"].ingest(
        payload,
        filename="depth.txt",
        content_type="text/plain",
        project_id=context["project"].id,
        well_id=context["well"].id,
    )
    assert outcome.document is not None
    record = (
        await session.execute(select(ExtractedRecord).where(ExtractedRecord.record_type == "depth_reading"))
    ).scalars().one()
    link = (
        await session.execute(select(EvidenceLink).where(EvidenceLink.subject_id == record.id))
    ).scalars().one()
    locator = link.locator["field_locations"]["bit_depth_si"]
    page = (
        await session.execute(select(DocumentPage).where(DocumentPage.document_id == outcome.document.id))
    ).scalars().one()
    assert page.text[locator["char_start"]:locator["char_end"]] == "Bit Depth: 1500 m"
    assert locator["line"] == 2
    assert locator["line_char_start"] == 2
    assert record.region_id is not None
    region = (await session.execute(select(DocumentRegion).where(DocumentRegion.id == record.region_id))).scalar_one()
    assert locator["region_order"] == region.order_index


async def test_scanned_document_records_ocr_requirement(session, context):
    outcome = await context["pipeline"].ingest(
        scanned_pdf(), filename="scan.pdf", content_type="application/pdf", well_id=context["well"].id
    )
    assert outcome.document is not None
    assert outcome.document.ocr_required is True
    assert any("OCR" in warning for warning in outcome.job.warnings)
    # A scanned PDF parses to a page with no text, so nothing could be extracted. That is an
    # incomplete read and the status now says so; the old values here (`parsed`, `ingested`) were
    # aliases for "finished" that could not distinguish it from a complete extraction.
    assert outcome.document.status == "partially_extracted"
    # The job still reports its own outcome separately: the run did what it was asked to do.
    assert outcome.job.status == "partial"


async def test_empty_file_is_refused(session, context):
    with pytest.raises(IngestionError):
        await context["pipeline"].ingest(b"", filename="empty.pdf")


async def test_failed_ingestion_can_commit_its_failure_record(session, context):
    """The API can retain a failed job/document even when its request later raises."""
    from drillai.core.errors import ExtractionFailed

    factory = async_sessionmaker(bind=session.bind, expire_on_commit=False)
    with pytest.raises(ExtractionFailed):
        await context["pipeline"].ingest(
            b"%PDF-1.4\nthis is not a real pdf",
            filename="broken.pdf",
            content_type="application/pdf",
            well_id=context["well"].id,
            commit=True,
        )
    async with factory() as verification_session:
        assert len((await verification_session.execute(select(RawArtifact))).scalars().all()) == 1
        job = (await verification_session.execute(select(IngestionJob))).scalars().one()
        assert job.status == "failed"
        assert job.error
        document = (await verification_session.execute(select(Document))).scalars().one()
        assert document.status == "failed"


async def test_blob_is_content_addressed(session, context):
    store = context["store"]
    outcome = await context["pipeline"].ingest(minimal_pdf(), filename="ddr.pdf")
    assert outcome.raw_artifact.blob_key == f"raw/{outcome.raw_artifact.sha256[:2]}/{outcome.raw_artifact.sha256[2:]}"
    assert store.exists(outcome.raw_artifact.blob_key)
    assert store.get(outcome.raw_artifact.blob_key) == minimal_pdf()
