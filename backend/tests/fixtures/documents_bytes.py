"""Byte-level document fixtures shared by the ingestion and document-domain suites.

These live outside the test modules because two suites need them: the ingestion pipeline tests build a
PDF, a DOCX and an XLSX to check that each parser is wired up, and the document-domain tests need the
same PDFs to check what the platform *records* about an ingestion. A second copy of a PDF assembler is
a second thing to keep valid.

Nothing here is generated at import time and nothing reaches the database: each function returns bytes
that are handed to the real parser.
"""

from __future__ import annotations

import io

# --------------------------------------------------------------------------- fixtures


def minimal_pdf(lines: list[str] | None = None) -> bytes:
    """A valid, uncompressed one-page PDF (pypdf extracts the text)."""
    body = lines or [
        "Daily Drilling Report",
        "Well: NF-12",
        "Rig: Rig 42",
        "Mud Weight: 9.2 ppg",
        "Plastic Viscosity: 18 cp",
        "Bit Depth: 1500 m",
        "Hole Depth: 1502 m",
    ]
    content = ("BT /F1 12 Tf 72 720 Td 14 TL\n" + "\n".join(f"({line}) Tj T*" for line in body) + "\nET").encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_position = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_position}\n%%EOF\n".encode()
    return bytes(out)


def scanned_pdf() -> bytes:
    """A PDF whose page has no text operators (stands in for a scan)."""
    content = b"0.5 0.5 0.5 rg 100 100 200 200 re f"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_position = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_position}\n%%EOF\n".encode()
    return bytes(out)


def docx_bytes() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Drilling Program - NF-12", level=1)
    document.add_paragraph("Well: NF-12")
    document.add_paragraph("Rig: Rig 42")
    table = document.add_table(rows=3, cols=4)
    header = ["MD (m)", "Inclination (deg)", "Azimuth (deg)", "Toolface"]
    for index, value in enumerate(header):
        table.rows[0].cells[index].text = value
    rows = [["0", "0.0", "0.0", "-"], ["500", "1.5", "45.0", "20"]]
    for row_index, values in enumerate(rows, start=1):
        for column, value in enumerate(values):
            table.rows[row_index].cells[column].text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes() -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Mud Check"
    rows = [
        ["Time", "Mud Weight (ppg)", "Funnel Viscosity (s)", "PV (cp)", "YP (lb/100ft2)"],
        ["06:00", 9.2, 45, 18, 12],
        ["12:00", 9.4, 48, 20, 14],
    ]
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
