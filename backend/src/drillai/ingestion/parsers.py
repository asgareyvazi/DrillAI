"""Format parsers: bytes in, a page/region model out.

Parsers are deliberately honest about what they can and cannot see:

* every parsed unit is a **page** with a **layout**;
* inside a page, content is expressed as **regions** — text blocks, tables and figures — each
  with an optional bounding box. Tables keep their grid (``rows``/``cols`` + ``cells`` with
  coordinates) because drilling data lives in tables and flattening them destroys meaning;
* page text is retained as well, because report readers (DDR header fields, narrative
  paragraphs) need the flow, and RAG needs something to index;
* nothing is invented: a parser that fails raises :class:`UnsupportedFormat` or
  :class:`ExtractionFailed` and the ingestion job records the failure.

Supported today (dependencies already in the project): PDF (``pypdf``), DOCX (``python-docx``),
XLSX/XLSM (``openpyxl``), CSV/TSV, and plain text/markdown. Scanned/image PDFs are detected
(no extractable text) and flagged ``ocr_required`` instead of silently returning an empty
document.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass, field
from typing import Any, BinaryIO, Protocol
from xml.etree import ElementTree
from xml.parsers import expat

from drillai.core.errors import ExtractionFailed, UnsupportedFormat

__all__ = [
    "PARSER_REGISTRY",
    "TABLE_HEADER_TOKENS",
    "ParsedDocument",
    "ParsedPage",
    "ParsedRegion",
    "Parser",
    "build_table",
    "detect_content_type",
    "detect_header_row",
    "parse_bytes",
    "parse_csv",
    "parse_docx",
    "parse_pdf",
    "parse_text",
    "parse_xlsx",
    "register_parser",
]

MAX_PAGES = 2000
MAX_BYTES = 64 * 1024 * 1024
MAX_SPREADSHEET_ROWS = 100_000
MAX_SPREADSHEET_COLUMNS = 2_000
MAX_SPREADSHEET_CELLS = 100_000
MAX_SPREADSHEET_RECORD_CHARS = 1_000_000
MAX_XLSX_UNCOMPRESSED_BYTES = 128 * 1024 * 1024
MAX_XLSX_ARCHIVE_MEMBERS = 10_000
MAX_DOCX_UNCOMPRESSED_BYTES = 128 * 1024 * 1024
MAX_DOCX_ARCHIVE_MEMBERS = 10_000


#: Column-name vocabulary used to locate the header row of a table. Real DDR/program exports from
#: rig reporting systems carry a title and metadata block *above* the operations table, so
#: "row 0 is the header" is wrong for the most important documents the platform ingests. The
#: detector scores candidate rows against these tokens and picks the best match; a table whose
#: first row already looks like a header keeps that row, so nothing regresses.
TABLE_HEADER_TOKENS: frozenset[str] = frozenset(
    {
        "operation", "operations", "activity", "activities", "description", "code", "duration",
        "start", "end", "from", "to", "depth", "md", "tvd", "inclination", "azimuth", "dls",
        "wob", "rpm", "flow", "flow rate", "spp", "torque", "hookload", "rop", "ecd", "mse",
        "bit", "bha", "nozzle", "tfa", "size", "od", "id", "weight", "grade", "pressure",
        "temperature", "mud", "mud weight", "pv", "yp", "date", "dates", "time", "times", "well",
        "section", "sections", "casing", "cement", "remarks", "comment", "comments", "qty", "quantity",
        "unit", "units", "value", "values", "parameter", "parameters", "survey", "station", "npt",
        "loss", "problem", "problems", "cost", "hours", "hrs",
    }
)

#: How far into a sheet/report the header row is searched. Deeper than this and the "table" is
#: really a narrative block.
MAX_HEADER_SEARCH_ROWS = 25


def _normalise_cell(value: str) -> str:
    return re.sub(r"[^a-z0-9 ]+", " ", str(value).lower()).strip()


def _contains_header_token(cell: str, token: str) -> bool:
    """Match a whole word/phrase, never an incidental substring inside another label."""
    normalised_cell = " ".join(cell.split())
    return f" {token} " in f" {normalised_cell} "


def detect_header_row(rows: list[list[str]]) -> tuple[int, list[str], float]:
    """Locate the header row of a table.

    Returns ``(index, header_values, confidence)``. Scoring is deliberately simple and
    explainable: a cell counts when it contains a known whole-word header token or phrase, and the
    best candidate wins on (number of token hits, number of non-empty cells). Rows that are mostly
    numbers are penalised, because a data row is not a header.
    """
    best_index, best_score = 0, -1.0
    for index, row in enumerate(rows[:MAX_HEADER_SEARCH_ROWS]):
        cells = [_normalise_cell(cell) for cell in row]
        non_empty = [cell for cell in cells if cell]
        if not non_empty:
            continue
        hits = sum(
            1
            for cell in non_empty
            if any(_contains_header_token(cell, token) for token in TABLE_HEADER_TOKENS)
        )
        numeric = sum(1 for cell in non_empty if re.fullmatch(r"-?\d+([.,]\d+)?", cell))
        if numeric == len(non_empty):
            continue
        score = hits * 2.0 - numeric * 0.5 + len(non_empty) * 0.05
        if score > best_score:
            best_index, best_score = index, score
    header = [str(cell) for cell in rows[best_index]] if rows else []
    hits = sum(
        1
        for cell in (_normalise_cell(cell) for cell in header)
        if cell and any(_contains_header_token(cell, token) for token in TABLE_HEADER_TOKENS)
    )
    confidence = min(1.0, 0.5 + 0.25 * hits) if hits else 0.4
    return best_index, header, confidence


def _excel_column_name(index: int) -> str:
    """Return the one-based column name used by spreadsheet locators."""
    letters = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def build_table(
    rows: list[list[str]],
    *,
    sheet: str | None = None,
    source_format: str = "delimited",
    source_rows: Sequence[int | dict[str, int] | None] | None = None,
    cell_metadata: list[list[dict[str, Any]]] | None = None,
    table_metadata: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Build a table while preserving original row/cell coordinates.

    ``cells`` retains blank rows and title blocks. ``row``/``col`` are zero-based offsets inside
    this serialized table; ``source_locator`` is the one-based position in the source file. The
    distinction matters for spreadsheets with deleted/blank rows and for CSV fields containing
    quoted newlines.
    """
    text_rows = [["" if value is None else str(value) for value in row] for row in rows]
    header_row, header, confidence = detect_header_row(text_rows)
    width = max((len(row) for row in text_rows), default=0)
    cells: list[list[dict[str, Any]]] = []
    for row_index, row in enumerate(text_rows):
        source = source_rows[row_index] if source_rows and row_index < len(source_rows) else row_index + 1
        if isinstance(source, dict):
            source_row = source.get("row")
            source_extra = {key: value for key, value in source.items() if key != "row"}
        else:
            source_row = source
            source_extra = {}
        cell_row: list[dict[str, Any]] = []
        for col_index, value in enumerate(row):
            locator: dict[str, Any] = {
                "format": source_format,
                "row": source_row,
                "column": col_index + 1,
                **source_extra,
            }
            if sheet is not None:
                locator["sheet"] = sheet
            if source_format == "xlsx":
                locator["cell"] = f"{_excel_column_name(col_index + 1)}{source_row}"
            cell = {
                "row": row_index,
                "col": col_index,
                "source_row": source_row,
                "source_column": col_index + 1,
                "source_locator": locator,
                "text": value,
            }
            if cell_metadata and row_index < len(cell_metadata) and col_index < len(cell_metadata[row_index]):
                cell.update(cell_metadata[row_index][col_index])
            cell_row.append(cell)
        cells.append(cell_row)
    table: dict[str, Any] = {
        "rows": len(text_rows),
        "cols": width,
        "cells": cells,
        "header": header,
        "header_row": header_row,
        "header_confidence": round(confidence, 3),
        "source_format": source_format,
    }
    if sheet is not None:
        table["sheet"] = sheet
    if table_metadata:
        table.update(table_metadata)
    return table


@dataclass
class ParsedRegion:
    kind: str  # text | table | figure | header_footer
    order_index: int
    text: str | None = None
    bbox: dict[str, float] | None = None  # {x0, y0, x1, y1} in page points
    table: dict[str, Any] | None = None  # {rows, cols, cells: [[{row, col, text}]], header: [...]}
    confidence: float | None = None
    source_locator: dict[str, Any] = field(default_factory=dict)


@dataclass
class ParsedPage:
    page_number: int
    text: str = ""
    width_pt: float | None = None
    height_pt: float | None = None
    regions: list[ParsedRegion] = field(default_factory=list)
    has_tables: bool = False
    has_figures: bool = False
    layout: dict[str, Any] = field(default_factory=dict)


@dataclass
class ParsedDocument:
    pages: list[ParsedPage]
    content_type: str
    parser: str
    parser_version: str
    metadata: dict[str, Any] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(page.text for page in self.pages if page.text)


class Parser(Protocol):
    content_types: tuple[str, ...]
    name: str
    version: str

    def parse(self, data: bytes, *, filename: str) -> ParsedDocument:  # pragma: no cover - protocol
        ...


PARSER_REGISTRY: dict[str, Parser] = {}


def register_parser(parser: Parser, *, replace: bool = False) -> None:
    for content_type in parser.content_types:
        if content_type in PARSER_REGISTRY and not replace:
            raise ValueError(f"parser for {content_type!r} already registered")
        PARSER_REGISTRY[content_type] = parser


_EXTENSION_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xlsm": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv",
    ".tsv": "text/tab-separated-values",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".log": "text/plain",
}


def detect_content_type(filename: str, declared: str | None = None) -> str:
    """Trust the declared type only if it is supported; otherwise infer from the extension."""
    if declared and declared in PARSER_REGISTRY:
        return declared
    lowered = filename.lower()
    for extension, content_type in _EXTENSION_TYPES.items():
        if lowered.endswith(extension):
            return content_type
    return declared or "application/octet-stream"


def parse_bytes(data: bytes, *, filename: str, content_type: str | None = None) -> ParsedDocument:
    if not data:
        raise ExtractionFailed("file is empty")
    if len(data) > MAX_BYTES:
        raise ExtractionFailed(f"file is larger than the {MAX_BYTES // (1024 * 1024)} MiB ingestion limit")
    resolved = detect_content_type(filename, content_type)
    parser = PARSER_REGISTRY.get(resolved)
    if parser is None:
        raise UnsupportedFormat(
            f"no parser registered for {resolved!r}",
            details={"filename": filename, "supported": sorted(PARSER_REGISTRY)},
        )
    return parser.parse(data, filename=filename)


# --------------------------------------------------------------------------- text


class TextParser:
    content_types = ("text/plain", "text/markdown")
    name = "text"
    version = "1.1.0"

    def parse(self, data: bytes, *, filename: str) -> ParsedDocument:
        text, encoding, decode_warnings = _decode_with_encoding(data)
        page = ParsedPage(
            page_number=1,
            text=text,
            regions=[
                ParsedRegion(
                    kind="text",
                    order_index=0,
                    text=text,
                    confidence=1.0,
                    source_locator={"format": "text", "page": 1},
                )
            ],
            layout={"source": "plain_text", "lines": text.count("\n") + 1, "encoding": encoding},
        )
        return ParsedDocument(
            pages=[page],
            content_type="text/plain",
            parser=self.name,
            parser_version=self.version,
            metadata={"filename": filename, "encoding": encoding},
            warnings=decode_warnings,
        )


def parse_text(data: bytes, *, filename: str) -> ParsedDocument:
    return TextParser().parse(data, filename=filename)


# --------------------------------------------------------------------------- csv / tsv


class _BoundedCsvLines:
    """Count physical input lines per logical CSV record before csv.reader materializes it."""

    def __init__(self, text: str, *, max_record_chars: int) -> None:
        self._lines = iter(io.StringIO(text, newline=""))
        self._max_record_chars = max_record_chars
        self._record_chars = 0

    def __iter__(self):
        return self

    def __next__(self) -> str:
        line = next(self._lines)
        self._record_chars += len(line)
        if self._record_chars > self._max_record_chars:
            raise ExtractionFailed(
                f"delimited record exceeds the {self._max_record_chars}-character processing limit"
            )
        return line

    def record_complete(self) -> None:
        self._record_chars = 0


class CsvParser:
    content_types = ("text/csv", "text/tab-separated-values")
    name = "csv"
    version = "1.1.0"

    @staticmethod
    def _delimiter(text: str, filename: str) -> str:
        if filename.lower().endswith(".tsv"):
            return "\t"
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            if dialect.delimiter in {",", ";", "\t"}:
                return dialect.delimiter
        except csv.Error:
            pass

        # The sniffer can reject short or uneven exports. Score actual parsed row widths rather
        # than counting punctuation, so commas inside quoted values do not masquerade as columns.
        scores: list[tuple[int, float, int, str]] = []
        for priority, delimiter in enumerate((",", ";", "\t")):
            try:
                parsed = list(csv.reader(io.StringIO(sample, newline=""), delimiter=delimiter, strict=True))[:30]
            except csv.Error:
                continue
            widths = [len(row) for row in parsed if row]
            if not widths:
                continue
            width = max(widths)
            consistency = sum(count == width for count in widths) / len(widths)
            scores.append((width, consistency, -priority, delimiter))
        if not scores:
            return ","
        best = max(scores)
        return best[3] if best[0] > 1 else ","

    def parse(self, data: bytes, *, filename: str) -> ParsedDocument:
        text, encoding, decode_warnings = _decode_with_encoding(data)
        delimiter = self._delimiter(text, filename)
        bounded_lines = _BoundedCsvLines(text, max_record_chars=MAX_SPREADSHEET_RECORD_CHARS)
        reader = csv.reader(bounded_lines, delimiter=delimiter, strict=True)
        rows: list[list[str]] = []
        source_rows: list[dict[str, int]] = []
        previous_end_line = 0
        cell_count = 0
        try:
            for row in reader:
                bounded_lines.record_complete()
                if len(rows) >= MAX_SPREADSHEET_ROWS:
                    raise ExtractionFailed(
                        f"delimited file exceeds the {MAX_SPREADSHEET_ROWS}-row processing limit"
                    )
                if len(row) > MAX_SPREADSHEET_COLUMNS:
                    raise ExtractionFailed(
                        f"delimited row exceeds the {MAX_SPREADSHEET_COLUMNS}-column processing limit"
                    )
                if cell_count + len(row) > MAX_SPREADSHEET_CELLS:
                    raise ExtractionFailed(
                        f"delimited file exceeds the {MAX_SPREADSHEET_CELLS}-cell processing limit"
                    )
                cell_count += len(row)
                start_line = previous_end_line + 1
                end_line = reader.line_num
                previous_end_line = end_line
                rows.append(row)
                source_rows.append({"row": start_line, "end_line": end_line})
        except csv.Error as exc:
            raise ExtractionFailed(f"malformed delimited text: {exc}") from exc
        if not rows:
            raise ExtractionFailed("CSV file has no rows")
        source_format = "tsv" if delimiter == "\t" else "csv"
        table = build_table(rows, source_format=source_format, source_rows=source_rows)
        header_width = len(table["header"])
        mismatched = [
            {"row": source_rows[index]["row"], "actual": len(row), "expected": header_width}
            for index, row in enumerate(rows[table["header_row"]:], start=table["header_row"])
            if len(row) != header_width
        ]
        warnings = list(decode_warnings)
        if mismatched:
            table["row_width_mismatches"] = mismatched[:100]
            warnings.append(
                f"{len(mismatched)} delimited row(s) differ from the detected header width; "
                "original columns were preserved without shifting values"
            )
        rendered = io.StringIO(newline="")
        csv.writer(rendered, delimiter=delimiter, lineterminator="\n").writerows(rows)
        page = ParsedPage(
            page_number=1,
            text=rendered.getvalue(),
            regions=[
                ParsedRegion(
                    kind="table",
                    order_index=0,
                    table=table,
                    confidence=1.0,
                    source_locator={"format": source_format, "delimiter": delimiter},
                )
            ],
            has_tables=True,
            layout={"source": "delimited", "delimiter": delimiter, "rows": len(rows), "encoding": encoding},
        )
        return ParsedDocument(
            pages=[page],
            content_type="text/tab-separated-values" if delimiter == "\t" else "text/csv",
            parser=self.name,
            parser_version=self.version,
            metadata={
                "filename": filename,
                "row_count": len(rows),
                "columns": table["header"],
                "header_row": table["header_row"],
                "encoding": encoding,
                "delimiter": delimiter,
            },
            warnings=warnings,
        )


def parse_csv(data: bytes, *, filename: str) -> ParsedDocument:
    return CsvParser().parse(data, filename=filename)


# --------------------------------------------------------------------------- docx


def _office_xml_parser(format_name: str, member_name: str) -> expat.XMLParserType:
    """Create an Expat parser that refuses DTDs before any declared entity can be expanded."""
    parser = expat.ParserCreate()

    def reject_dtd(*_args: Any) -> None:
        raise ExtractionFailed(
            f"{format_name} archive member {member_name!r} contains a forbidden DTD/entity declaration"
        )

    parser.StartDoctypeDeclHandler = reject_dtd
    parser.ExternalEntityRefHandler = reject_dtd
    parser.SetParamEntityParsing(expat.XML_PARAM_ENTITY_PARSING_NEVER)
    return parser


def _validate_office_archive(
    data: bytes,
    *,
    format_name: str,
    max_members: int,
    max_uncompressed_bytes: int,
) -> None:
    """Bound declared and actual ZIP expansion before parsing OOXML package parts."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > max_members:
                raise ExtractionFailed(
                    f"{format_name} archive exceeds the {max_members}-member processing limit"
                )
            declared_size = sum(member.file_size for member in members)
            if declared_size > max_uncompressed_bytes:
                raise ExtractionFailed(
                    f"{format_name} archive exceeds the "
                    f"{max_uncompressed_bytes // (1024 * 1024)} MiB uncompressed processing limit"
                )

            actual_size = 0
            for member in members:
                if member.is_dir():
                    continue
                member_size = 0
                is_xml = member.filename.casefold().endswith((".xml", ".rels", ".vml"))
                xml_parser = _office_xml_parser(format_name, member.filename) if is_xml else None
                with archive.open(member, "r") as stream:
                    while block := stream.read(64 * 1024):
                        member_size += len(block)
                        actual_size += len(block)
                        if actual_size > max_uncompressed_bytes:
                            raise ExtractionFailed(
                                f"{format_name} archive exceeds the "
                                f"{max_uncompressed_bytes // (1024 * 1024)} MiB uncompressed processing limit"
                            )
                        if xml_parser is not None:
                            xml_parser.Parse(block, False)
                    if xml_parser is not None:
                        xml_parser.Parse(b"", True)
                if member_size != member.file_size:
                    raise ExtractionFailed(
                        f"{format_name} archive member {member.filename!r} has an invalid uncompressed size"
                    )
    except ExtractionFailed:
        raise
    except Exception as exc:
        raise ExtractionFailed(f"cannot read {format_name}: invalid or unsafe ZIP archive ({exc})") from exc


def _validate_docx_archive(data: bytes) -> None:
    _validate_office_archive(
        data,
        format_name="DOCX",
        max_members=MAX_DOCX_ARCHIVE_MEMBERS,
        max_uncompressed_bytes=MAX_DOCX_UNCOMPRESSED_BYTES,
    )


class DocxParser:
    content_types = ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",)
    name = "docx"
    version = "1.1.0"

    def parse(self, data: bytes, *, filename: str) -> ParsedDocument:
        import docx  # python-docx
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        _validate_docx_archive(data)
        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:
            raise ExtractionFailed(f"cannot read DOCX: {exc}") from exc

        regions: list[ParsedRegion] = []
        text_parts: list[str] = []
        text_length = 0

        def append_text_part(text: str) -> tuple[int, int]:
            nonlocal text_length
            if text_parts:
                text_length += 1  # The page text joins every logical block with one newline.
            start = text_length
            text_parts.append(text)
            text_length += len(text)
            return start, text_length

        order = 0
        paragraph_number = 0
        table_number = 0
        for block in document.element.body.iterchildren():
            if isinstance(block, CT_P):
                paragraph_number += 1
                paragraph = Paragraph(block, document).text.strip()
                if paragraph:
                    char_start, char_end = append_text_part(paragraph)
                    regions.append(
                        ParsedRegion(
                            kind="text",
                            order_index=order,
                            text=paragraph,
                            confidence=1.0,
                            source_locator={
                                "format": "docx",
                                "paragraph": paragraph_number,
                                "char_start": char_start,
                                "char_end": char_end,
                            },
                        )
                    )
                    order += 1
            elif isinstance(block, CT_Tbl):
                table_number += 1
                source_table = Table(block, document)
                rows: list[list[str]] = []
                cell_metadata: list[list[dict[str, Any]]] = []
                merged_repeats = 0
                for row in source_table.rows:
                    values: list[str] = []
                    metadata_row: list[dict[str, Any]] = []
                    seen_cells: set[int] = set()
                    for cell in row.cells:
                        identity = id(cell._tc)
                        repeated = identity in seen_cells
                        seen_cells.add(identity)
                        actual_text = cell.text.strip()
                        values.append("" if repeated else actual_text)
                        merged_repeats += int(repeated)
                        metadata_row.append(
                            {
                                "value_type": "text",
                                "display_text": actual_text,
                                "merged_repeat": repeated,
                            }
                        )
                    rows.append(values)
                    cell_metadata.append(metadata_row)
                table_payload = build_table(
                    rows,
                    source_format="docx",
                    source_rows=[{"row": index, "table": table_number} for index in range(1, len(rows) + 1)],
                    cell_metadata=cell_metadata,
                    table_metadata={"table_number": table_number, "merged_repeat_count": merged_repeats},
                )
                regions.append(
                    ParsedRegion(
                        kind="table",
                        order_index=order,
                        table=table_payload,
                        confidence=1.0,
                        source_locator={"format": "docx", "table": table_number},
                    )
                )
                for row in table_payload["cells"]:
                    append_text_part(" | ".join(cell.get("display_text", cell["text"]) for cell in row))
                order += 1
        if not regions:
            raise ExtractionFailed("DOCX contains no readable text or tables")

        page = ParsedPage(
            page_number=1,
            text="\n".join(text_parts),
            regions=regions,
            has_tables=any(region.kind == "table" for region in regions),
            layout={
                "source": "docx",
                "tables": table_number,
                "paragraphs": paragraph_number,
                "page_number_is_logical": True,
            },
        )
        return ParsedDocument(
            pages=[page],
            content_type=self.content_types[0],
            parser=self.name,
            parser_version=self.version,
            metadata={"filename": filename, "page_count_is_logical": True},
        )


def parse_docx(data: bytes, *, filename: str) -> ParsedDocument:
    return DocxParser().parse(data, filename=filename)


# --------------------------------------------------------------------------- xlsx


def _xlsx_coordinate(reference: str) -> tuple[int, int]:
    match = re.fullmatch(r"\$?([A-Z]{1,3})\$?([1-9][0-9]*)", reference, re.I)
    if match is None:
        raise ExtractionFailed(f"XLSX worksheet contains an invalid cell reference {reference!r}")
    column = 0
    for letter in match.group(1).upper():
        column = column * 26 + ord(letter) - ord("A") + 1
    return int(match.group(2)), column


def _validate_xlsx_dimensions(data: bytes) -> None:
    """Reject oversized worksheet dimensions while streaming XML, before openpyxl builds cells."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            worksheets = sorted(
                name
                for name in archive.namelist()
                if name.startswith("xl/worksheets/")
                and name.endswith(".xml")
                and "/_rels/" not in name
            )
            if len(worksheets) > MAX_PAGES:
                raise ExtractionFailed(f"XLSX workbook exceeds the {MAX_PAGES}-worksheet processing limit")

            total_area = 0
            for worksheet_name in worksheets:
                max_row = 0
                max_column = 0
                with archive.open(worksheet_name) as stream:
                    for event, element in ElementTree.iterparse(stream, events=("start", "end")):
                        if event == "start":
                            tag = element.tag.rsplit("}", 1)[-1]
                            if tag == "dimension":
                                reference = element.attrib.get("ref", "")
                                endpoints = reference.split(":")
                                if not endpoints or len(endpoints) > 2:
                                    raise ExtractionFailed(
                                        f"XLSX worksheet {worksheet_name!r} has an invalid dimension {reference!r}"
                                    )
                                coordinates = [_xlsx_coordinate(endpoint) for endpoint in endpoints]
                                max_row = max(max_row, *(row for row, _column in coordinates))
                                max_column = max(max_column, *(column for _row, column in coordinates))
                            elif tag == "row":
                                row_ref = element.attrib.get("r")
                                if row_ref is not None:
                                    try:
                                        max_row = max(max_row, int(row_ref))
                                    except ValueError as exc:
                                        raise ExtractionFailed(
                                            f"XLSX worksheet {worksheet_name!r} has an invalid row index"
                                        ) from exc
                            elif tag == "c":
                                cell_ref = element.attrib.get("r")
                                if cell_ref is not None:
                                    row, column = _xlsx_coordinate(cell_ref)
                                    max_row = max(max_row, row)
                                    max_column = max(max_column, column)
                            if max_row > MAX_SPREADSHEET_ROWS:
                                raise ExtractionFailed(
                                    f"worksheet {worksheet_name!r} exceeds the supported row limit ({max_row} rows)"
                                )
                            if max_column > MAX_SPREADSHEET_COLUMNS:
                                raise ExtractionFailed(
                                    f"worksheet {worksheet_name!r} exceeds the supported column limit "
                                    f"({max_column} columns)"
                                )
                            if total_area + max_row * max_column > MAX_SPREADSHEET_CELLS:
                                raise ExtractionFailed(
                                    f"workbook exceeds the {MAX_SPREADSHEET_CELLS}-cell processing limit"
                                )
                        elif element.tag.rsplit("}", 1)[-1] in {"c", "row"}:
                            element.clear()
                total_area += max_row * max_column
    except ExtractionFailed:
        raise
    except Exception as exc:
        raise ExtractionFailed(f"cannot inspect XLSX worksheet dimensions: {exc}") from exc


def _validate_xlsx_archive(data: bytes) -> None:
    """Bound ZIP expansion and cell dimensions before openpyxl materialises workbook objects."""
    _validate_office_archive(
        data,
        format_name="XLSX",
        max_members=MAX_XLSX_ARCHIVE_MEMBERS,
        max_uncompressed_bytes=MAX_XLSX_UNCOMPRESSED_BYTES,
    )
    _validate_xlsx_dimensions(data)


class XlsxParser:
    content_types = ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",)
    name = "xlsx"
    version = "1.1.0"

    @staticmethod
    def _value_type(value: Any) -> str:
        if value is None:
            return "blank"
        if isinstance(value, bool):
            return "boolean"
        if isinstance(value, (int, float)):
            return "number"
        if isinstance(value, dt.datetime):
            return "datetime"
        if isinstance(value, dt.date):
            return "date"
        if isinstance(value, dt.time):
            return "time"
        return "text"

    @staticmethod
    def _json_value(value: Any) -> Any:
        if isinstance(value, (dt.datetime, dt.date, dt.time)):
            return value.isoformat()
        if isinstance(value, (str, int, float, bool)) or value is None:
            return value
        return str(value)

    @staticmethod
    def _text_value(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, dt.datetime):
            return value.isoformat(sep=" ")
        if isinstance(value, (dt.date, dt.time)):
            return value.isoformat()
        if isinstance(value, bool):
            return "TRUE" if value else "FALSE"
        return str(value)

    def parse(self, data: bytes, *, filename: str) -> ParsedDocument:
        import openpyxl

        _validate_xlsx_archive(data)
        formula_workbook = None
        value_workbook = None
        try:
            formula_workbook = openpyxl.load_workbook(
                io.BytesIO(data), data_only=False, read_only=False, keep_links=False
            )
            value_workbook = openpyxl.load_workbook(
                io.BytesIO(data), data_only=True, read_only=False, keep_links=False
            )
        except Exception as exc:
            for workbook in (formula_workbook, value_workbook):
                if workbook is not None:
                    workbook.close()
            raise ExtractionFailed(f"cannot read XLSX: {exc}") from exc

        pages: list[ParsedPage] = []
        warnings: list[str] = []
        formula_total = 0
        hidden_sheet_count = 0
        hidden_row_total = 0
        total_cells = 0
        try:
            for index, sheet in enumerate(formula_workbook.worksheets[:MAX_PAGES], start=1):
                value_sheet = value_workbook[sheet.title]
                row_count = int(sheet.max_row or 0)
                column_count = int(sheet.max_column or 0)
                if row_count > MAX_SPREADSHEET_ROWS or column_count > MAX_SPREADSHEET_COLUMNS:
                    raise ExtractionFailed(
                        f"worksheet {sheet.title!r} exceeds the supported row or column limit "
                        f"({row_count} rows, {column_count} columns)"
                    )
                sheet_cells = row_count * column_count
                if total_cells + sheet_cells > MAX_SPREADSHEET_CELLS:
                    raise ExtractionFailed(
                        f"workbook exceeds the {MAX_SPREADSHEET_CELLS}-cell processing limit"
                    )
                total_cells += sheet_cells
                if row_count == 0 or column_count == 0:
                    warnings.append(f"worksheet {sheet.title!r} is empty and was skipped")
                    continue

                hidden_sheet = sheet.sheet_state != "visible"
                hidden_rows = {
                    row_index
                    for row_index in range(1, row_count + 1)
                    if bool(sheet.row_dimensions[row_index].hidden)
                }
                hidden_row_total += len(hidden_rows)
                hidden_sheet_count += int(hidden_sheet)
                rows: list[list[str]] = []
                source_rows: list[int] = []
                cell_metadata: list[list[dict[str, Any]]] = []
                rendered_rows: list[str] = []
                formula_cells: list[str] = []
                sheet_has_content = False
                for row_index in range(1, row_count + 1):
                    values: list[str] = []
                    metadata_row: list[dict[str, Any]] = []
                    display_values: list[str] = []
                    row_hidden = row_index in hidden_rows
                    for column_index in range(1, column_count + 1):
                        cell = sheet.cell(row=row_index, column=column_index)
                        cached_cell = value_sheet.cell(row=row_index, column=column_index)
                        raw_value = cell.value
                        cached_value = cached_cell.value
                        is_formula = cell.data_type == "f"
                        sheet_has_content = sheet_has_content or raw_value is not None or cached_value is not None
                        if is_formula:
                            formula_total += 1
                            cell_ref = cell.coordinate
                            if len(formula_cells) < 100:
                                formula_cells.append(cell_ref)
                            extract_text = ""
                            visible_text = "[formula cell excluded from structured extraction]"
                            metadata = {
                                "value_type": "formula",
                                "formula": str(raw_value),
                                "cached_value": self._json_value(cached_value),
                                "cached_value_type": self._value_type(cached_value),
                                "cache_available": cached_value is not None,
                                "number_format": cell.number_format,
                                "display_text": visible_text,
                            }
                        else:
                            actual_text = self._text_value(raw_value)
                            extract_text = actual_text
                            visible_text = actual_text
                            metadata = {
                                "value_type": self._value_type(raw_value),
                                "value": self._json_value(raw_value),
                                "number_format": cell.number_format,
                                "display_text": actual_text,
                            }
                        if row_hidden or hidden_sheet:
                            extract_text = ""
                        values.append(extract_text)
                        display_values.append(visible_text)
                        metadata_row.append(
                            {
                                **metadata,
                                "formula_ignored": is_formula,
                                "hidden_row": row_hidden,
                                "hidden_sheet": hidden_sheet,
                            }
                        )
                    rows.append(values)
                    source_rows.append(row_index)
                    cell_metadata.append(metadata_row)
                    if hidden_sheet:
                        rendered_rows.append("[hidden worksheet content excluded from extraction]")
                    elif row_hidden:
                        rendered_rows.append("[hidden row content excluded from extraction]")
                    else:
                        rendered_rows.append(" | ".join(display_values))

                if not sheet_has_content:
                    warnings.append(f"worksheet {sheet.title!r} is empty and was skipped")
                    continue
                merged_ranges = sorted(str(item) for item in sheet.merged_cells.ranges)
                table = build_table(
                    rows,
                    sheet=sheet.title,
                    source_format="xlsx",
                    source_rows=source_rows,
                    cell_metadata=cell_metadata,
                    table_metadata={
                        "hidden": hidden_sheet,
                        "sheet_state": sheet.sheet_state,
                        "hidden_rows": sorted(hidden_rows),
                        "merged_ranges": merged_ranges,
                        "formula_cells": formula_cells,
                        "formula_cell_count": sum(
                            1 for row in cell_metadata for cell_meta in row if cell_meta.get("formula_ignored")
                        ),
                    },
                )
                page_text = f"# sheet: {sheet.title}"
                if hidden_sheet:
                    page_text += " [hidden]"
                page_text += "\n" + "\n".join(rendered_rows)
                pages.append(
                    ParsedPage(
                        page_number=index,
                        text=page_text,
                        regions=[
                            ParsedRegion(
                                kind="table",
                                order_index=0,
                                table=table,
                                confidence=1.0,
                                source_locator={"format": "xlsx", "sheet": sheet.title},
                            )
                        ],
                        has_tables=True,
                        layout={
                            "source": "xlsx",
                            "sheet": sheet.title,
                            "sheet_state": sheet.sheet_state,
                            "rows": row_count,
                            "columns": column_count,
                            "hidden_rows": sorted(hidden_rows),
                            "merged_ranges": merged_ranges,
                        },
                    )
                )
                if hidden_sheet:
                    warnings.append(f"worksheet {sheet.title!r} is hidden and excluded from extraction")
                if hidden_rows:
                    warnings.append(
                        f"worksheet {sheet.title!r} contains {len(hidden_rows)} hidden row(s); "
                        "hidden values are preserved as metadata but excluded from structured extraction"
                    )
                formula_count = table["formula_cell_count"]
                if formula_count:
                    warnings.append(
                        f"worksheet {sheet.title!r} contains {formula_count} formula cell(s); "
                        "formula text and cached values are stored separately and excluded from structured extraction"
                    )
                if merged_ranges:
                    warnings.append(
                        f"worksheet {sheet.title!r} contains {len(merged_ranges)} merged range(s); "
                        "merged anchor/value positions are preserved without filling blank cells"
                    )
        finally:
            formula_workbook.close()
            value_workbook.close()

        if not pages:
            raise ExtractionFailed("workbook contains no non-empty worksheets")
        return ParsedDocument(
            pages=pages,
            content_type=self.content_types[0],
            parser=self.name,
            parser_version=self.version,
            metadata={
                "filename": filename,
                "sheets": len(pages),
                "formula_cells": formula_total,
                "hidden_sheets": hidden_sheet_count,
                "hidden_rows": hidden_row_total,
            },
            warnings=warnings,
        )


def parse_xlsx(data: bytes, *, filename: str) -> ParsedDocument:
    return XlsxParser().parse(data, filename=filename)


# --------------------------------------------------------------------------- pdf


class PdfParser:
    content_types = ("application/pdf",)
    name = "pdf"
    version = "1.0.0"

    def parse(self, data: bytes, *, filename: str) -> ParsedDocument:
        import pypdf

        try:
            reader = pypdf.PdfReader(io.BytesIO(data))
        except Exception as exc:
            raise ExtractionFailed(f"cannot read PDF: {exc}") from exc
        if not reader.pages:
            raise ExtractionFailed("PDF contains no pages")

        pages: list[ParsedPage] = []
        warnings: list[str] = []
        for index, page in enumerate(reader.pages[:MAX_PAGES], start=1):
            try:
                text = page.extract_text() or ""
            except Exception as exc:
                text = ""
                warnings.append(f"page {index}: text extraction failed ({exc})")
            box = page.mediabox
            regions: list[ParsedRegion] = []
            if text.strip():
                regions.append(
                    ParsedRegion(
                        kind="text",
                        order_index=0,
                        text=text,
                        confidence=0.9,
                        source_locator={"format": "pdf", "page": index, "granularity": "page_text"},
                    )
                )
            pages.append(
                ParsedPage(
                    page_number=index,
                    text=text,
                    width_pt=float(box.width) if box else None,
                    height_pt=float(box.height) if box else None,
                    regions=regions,
                    layout={"source": "pdf", "extractor": "pypdf"},
                )
            )
        if len(reader.pages) > MAX_PAGES:
            warnings.append(f"document truncated to the first {MAX_PAGES} pages")
        ocr_required_pages = [page.page_number for page in pages if not page.text.strip()]
        if ocr_required_pages:
            if len(ocr_required_pages) == len(pages):
                warnings.append(
                    "no extractable text: the PDF is probably scanned (OCR required before extraction)"
                )
            else:
                warnings.append(
                    f"{len(ocr_required_pages)} PDF page(s) have no extractable text; OCR is required for those pages"
                )
        return ParsedDocument(
            pages=pages,
            content_type="application/pdf",
            parser=self.name,
            parser_version=self.version,
            metadata={
                "filename": filename,
                "page_count": len(pages),
                "ocr_required": bool(ocr_required_pages),
                "ocr_required_pages": ocr_required_pages,
            },
            warnings=warnings,
        )


def parse_pdf(data: bytes, *, filename: str) -> ParsedDocument:
    return PdfParser().parse(data, filename=filename)


# --------------------------------------------------------------------------- helpers


def _decode_with_encoding(data: bytes) -> tuple[str, str, list[str]]:
    """Decode common report exports without silently replacing undecodable bytes."""
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return data.decode("utf-16"), "utf-16", []
        except UnicodeDecodeError as exc:
            raise ExtractionFailed(f"invalid UTF-16 text: {exc}") from exc
    if data.startswith(b"\xef\xbb\xbf"):
        try:
            return data.decode("utf-8-sig"), "utf-8-sig", []
        except UnicodeDecodeError as exc:
            raise ExtractionFailed(f"invalid UTF-8 BOM text: {exc}") from exc
    try:
        return data.decode("utf-8"), "utf-8", []
    except UnicodeDecodeError:
        # Western rig-report exports commonly use Windows-1252. It differs from Latin-1 in the
        # 0x80-0x9f range, so try it first; Latin-1 remains a lossless final fallback.
        try:
            return data.decode("cp1252"), "cp1252", ["input was decoded as Windows-1252, not UTF-8"]
        except UnicodeDecodeError:
            return data.decode("latin-1"), "latin-1", ["input was decoded as Latin-1, not UTF-8"]


def _decode(data: bytes) -> str:
    return _decode_with_encoding(data)[0]


def looks_like_scanned(document: ParsedDocument) -> bool:
    return bool(document.metadata.get("ocr_required"))


_DRILLING_HINTS = (
    "daily drilling report",
    "drilling report",
    "mud weight",
    "bit depth",
    "hole depth",
    "survey",
    "casing",
)


def guess_doc_type(document: ParsedDocument, filename: str) -> str:
    """Cheap deterministic classification used to pick extractors (LLM-free by design)."""
    haystack = f"{filename}\n{document.text[:6000]}".lower()
    if re.search(r"\bdaily (drilling )?report\b", haystack) or "ddr" in filename.lower():
        return "ddr"
    if "end of well" in haystack or "eowr" in filename.lower():
        return "eowr"
    if "drilling program" in haystack or "well program" in haystack or "programme" in haystack:
        return "program"
    if "procedure" in haystack:
        return "procedure"
    if any(hint in haystack for hint in _DRILLING_HINTS):
        return "drilling_document"
    return "other"


for _parser in (TextParser(), CsvParser(), DocxParser(), XlsxParser(), PdfParser()):
    register_parser(_parser)


def parser_catalogue() -> list[dict[str, Any]]:
    return [
        {"content_type": content_type, "parser": parser.name, "version": parser.version}
        for content_type, parser in sorted(PARSER_REGISTRY.items())
    ]


def stream_to_bytes(stream: BinaryIO) -> bytes:  # pragma: no cover - convenience
    return stream.read()
